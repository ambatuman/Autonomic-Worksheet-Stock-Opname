import io
import openpyxl
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="Stock Opname Worksheet Generator", page_icon="📊", layout="wide"
)

st.title("📊 Stock Opname Worksheet Generator")
st.write(
    "Aplikasi pemroses & perapi dokumen Excel Stock Opname secara otomatis."
)

st.divider()

# ---------------------------------------------------------
# STEP 1: UPLOAD DOKUMEN
# ---------------------------------------------------------
st.header("1. Upload Dokumen")

col1, col2, col3 = st.columns(3)

with col1:
    template_file = st.file_uploader(
        "Upload Template Master (.xlsx)",
        type=["xlsx"],
        help="Dokumen master/template yang berisi sheet Worksheet, Summary, dll.",
    )

with col2:
    csv_file = st.file_uploader(
        "Upload Data Mentah Inventory (.csv / .xlsx)",
        type=["csv", "xlsx", "xls"],
        help="Report inventory baru (bisa format CSV atau Excel).",
    )

with col3:
    prev_so_file = st.file_uploader(
        "Upload Dokumen Prev SO Referensi (.xlsx)",
        type=["xlsx"],
        help="File SO periode sebelumnya untuk VLOOKUP Batch -> Prev SO & Prev Status.",
    )

st.divider()

# ---------------------------------------------------------
# STEP 2: INPUT METADATA & SUMMARY
# ---------------------------------------------------------
st.header("2. Input Informasi Metadata & Summary")

col_m1, col_m2, col_m3 = st.columns(3)
with col_m1:
    station_input = st.text_input("Station (Cell C3 / D3)", value="BRJ")
with col_m2:
    location_input = st.text_input(
        "Location (Cell C4)", value="BL04"
    )
with col_m3:
    periode_input = st.text_input(
        "Periode Date (Cell C5 / D4)", value="17 - 18 SEPTEMBER 2026"
    )

st.subheader("Pengaturan Sheet Summary (Dinamis Per Lokasi)")
st.caption(
    "Masukkan detail divisi, PIC, kode lokasi, dan deskripsi lokasi. Anda bisa menambah atau menghapus lokasi sesuai kebutuhan."
)

if "summary_rows" not in st.session_state:
    st.session_state.summary_rows = [
        {
            "division": "LINE MAINTENANCE",
            "pic": "THOMAS",
            "loc_code": "BL04",
            "loc_desc": "STORE BRJ",
        }
    ]

# Tombol Tambah / Hapus Lokasi
col_b1, col_b2, _ = st.columns([1, 1, 4])
with col_b1:
    if st.button("➕ Tambah Lokasi"):
        st.session_state.summary_rows.append(
            {"division": "", "pic": "", "loc_code": "", "loc_desc": ""}
        )
        st.rerun()

with col_b2:
    if (
        st.button("➖ Hapus Lokasi Terakhir")
        and len(st.session_state.summary_rows) > 1
    ):
        st.session_state.summary_rows.pop()
        st.rerun()

# Form dinamis untuk input summary
summary_data_inputs = []
for idx, row in enumerate(st.session_state.summary_rows):
    st.markdown(f"**Lokasi #{idx + 1}**")
    c1, c2, c3, c4 = st.columns(4)
    with c1:
        div = st.text_input(
            f"Division #{idx+1}", value=row["division"], key=f"div_{idx}"
        )
    with c2:
        pic = st.text_input(
            f"PIC #{idx+1}", value=row["pic"], key=f"pic_{idx}"
        )
    with c3:
        loc = st.text_input(
            f"LOC Code #{idx+1}", value=row["loc_code"], key=f"loc_{idx}"
        )
    with c4:
        desc = st.text_input(
            f"LOC Description #{idx+1}",
            value=row["loc_desc"],
            key=f"desc_{idx}",
        )
    summary_data_inputs.append(
        {"division": div, "pic": pic, "loc_code": loc, "loc_desc": desc}
    )

st.divider()

# ---------------------------------------------------------
# STEP 3: HELPER FUNCTIONS & EXEKUSI
# ---------------------------------------------------------
st.header("3. Eksekusi & Pemrosesan")


def load_raw_inventory_data(uploaded_file_obj):
    """Membaca file data mentah baik berupa Excel (.xlsx/.xls) maupun CSV."""
    fname = uploaded_file_obj.name.lower()
    if fname.endswith(".xlsx") or fname.endswith(".xls"):
        df = pd.read_excel(uploaded_file_obj)
    else:
        try:
            df = pd.read_csv(uploaded_file_obj, sep=";")
            if df.shape[1] <= 1:
                uploaded_file_obj.seek(0)
                df = pd.read_csv(uploaded_file_obj, sep=",")
        except Exception:
            uploaded_file_obj.seek(0)
            df = pd.read_csv(uploaded_file_obj, sep=",")

    # Clean & normalize column names to UPPERCASE
    df.columns = df.columns.astype(str).str.strip().str.upper()
    return df


def get_val(row, *possible_keys, default=None):
    """Helper untuk mengambil data secara Case-Insensitive dari dictionary baris."""
    for key in possible_keys:
        k_upper = key.strip().upper()
        if k_upper in row and pd.notna(row[k_upper]):
            return row[k_upper]
    return default


def set_cell_safe(ws, row, col, value):
    """Mengisi sel tanpa merusak/error jika menimpa MergedCell."""
    cell = ws.cell(row=row, column=col)
    if type(cell).__name__ == "MergedCell":
        # Jika sel bagian dari merged range, cari sel utama di pojok kiri atas
        for merged_range in ws.merged_cells.ranges:
            if cell.coordinate in merged_range:
                ws.cell(row=merged_range.min_row, column=merged_range.min_col).value = value
                break
    else:
        cell.value = value


def build_prev_so_dict(prev_so_file_obj):
    """Membaca file Prev SO dan membuat kamus lookup berdasarkan Batch Key."""
    prev_wb = openpyxl.load_workbook(prev_so_file_obj, data_only=True)
    prev_dict = {}

    if "Worksheet" in prev_wb.sheetnames:
        ws_prev = prev_wb["Worksheet"]

        batch_col, result_col, status_col = None, None, None
        header_row = 7

        for r in range(1, 10):
            for c in range(1, 40):
                val = str(ws_prev.cell(r, c).value or "").strip().lower()
                if val == "batch":
                    batch_col = c
                elif val == "result":
                    result_col = c
                elif val == "status":
                    status_col = c

            if batch_col and result_col and status_col:
                header_row = r
                break

        if batch_col:
            for r in range(header_row + 1, ws_prev.max_row + 1):
                batch_val = ws_prev.cell(r, batch_col).value
                if batch_val is not None and str(batch_val).strip() != "":
                    batch_key = str(batch_val).strip()
                    res_val = (
                        ws_prev.cell(r, result_col).value
                        if result_col
                        else None
                    )
                    stat_val = (
                        ws_prev.cell(r, status_col).value
                        if status_col
                        else None
                    )
                    prev_dict[batch_key] = {
                        "prev_so": res_val,
                        "prev_status": stat_val,
                    }
    return prev_dict


if st.button("🚀 Process & Generate Template", type="primary"):
    if not template_file or not csv_file or not prev_so_file:
        st.error(
            "⚠️ Harap upload KETIGA dokumen terlebih dahulu sebelum memproses!"
        )
    else:
        with st.spinner("Sedang memproses data dan merapikan Excel..."):
            wb = openpyxl.load_workbook(template_file)
            df_raw = load_raw_inventory_data(csv_file)
            prev_so_map = build_prev_so_dict(prev_so_file)

            # Hapus Sheet1 & Sheet2 jika ada
            for sname in ["Sheet1", "Sheet2", "sheet1", "sheet2"]:
                if sname in wb.sheetnames:
                    del wb[sname]

            # -----------------------------------------------------
            # UPDATE SHEET WORKSHEET
            # -----------------------------------------------------
            if "Worksheet" in wb.sheetnames:
                ws = wb["Worksheet"]

                # Update Metadata
                set_cell_safe(ws, 3, 3, f": {station_input}")
                set_cell_safe(ws, 4, 3, f": {location_input}")
                set_cell_safe(ws, 5, 3, f": {periode_input}")

                # Hapus data lama mulai baris 8 ke bawah
                max_r = ws.max_row
                if max_r >= 8:
                    ws.delete_rows(8, amount=max_r - 7 + 1)

                # Isi data mentah baru mulai dari baris ke-8
                for idx, row_data in df_raw.iterrows():
                    row_idx = 8 + idx

                    batch_num = str(
                        get_val(
                            row_data, "BATCH", "BATCH NO", "BATCH_NO", default=""
                        )
                    ).strip()

                    set_cell_safe(ws, row_idx, 1, idx + 1)  # A: No
                    set_cell_safe(
                        ws,
                        row_idx,
                        2,
                        get_val(row_data, "COUNT NO", "COUNT_NO", default=""),
                    )  # B: Count No
                    set_cell_safe(
                        ws,
                        row_idx,
                        3,
                        get_val(row_data, "LOCATION", "LOC", default=""),
                    )  # C: LOC
                    set_cell_safe(
                        ws, row_idx, 4, get_val(row_data, "BIN", default="")
                    )  # D: BIN
                    set_cell_safe(
                        ws,
                        row_idx,
                        5,
                        get_val(row_data, "GRB", "GRB NO", "GRB_NO", default=""),
                    )  # E: GRB
                    set_cell_safe(ws, row_idx, 6, batch_num)  # F: Batch
                    set_cell_safe(
                        ws,
                        row_idx,
                        7,
                        get_val(
                            row_data, "PN", "PART NO", "PART_NO", default=""
                        ),
                    )  # G: PN
                    set_cell_safe(
                        ws,
                        row_idx,
                        8,
                        get_val(
                            row_data, "SN", "SERIAL NO", "SERIAL_NO", default=""
                        ),
                    )  # H: SN
                    set_cell_safe(
                        ws,
                        row_idx,
                        9,
                        get_val(
                            row_data,
                            "PN DESCRIPTION",
                            "DESCRIPTION",
                            "PN DESC",
                            default="",
                        ),
                    )  # I: PN Description

                    # QTY Columns (10-15)
                    set_cell_safe(
                        ws,
                        row_idx,
                        10,
                        get_val(
                            row_data,
                            "QTY AVAILABLE",
                            "QTY_AVAIL",
                            "QTY AVAIL",
                            default=0,
                        ),
                    )  # J
                    set_cell_safe(
                        ws,
                        row_idx,
                        11,
                        get_val(
                            row_data,
                            "QTY RESERVED",
                            "QTY_RESV",
                            "QTY RESV",
                            default=0,
                        ),
                    )  # K
                    set_cell_safe(
                        ws,
                        row_idx,
                        12,
                        get_val(
                            row_data,
                            "QTY IN TRANSFER",
                            "QTY_TRANS",
                            "QTY TRANS",
                            default=0,
                        ),
                    )  # L
                    set_cell_safe(
                        ws,
                        row_idx,
                        13,
                        get_val(
                            row_data,
                            "QTY PENDING RI",
                            "QTY PENDING R/I",
                            "QTY_PENDING",
                            default=0,
                        ),
                    )  # M
                    set_cell_safe(
                        ws,
                        row_idx,
                        14,
                        get_val(row_data, "QTY US", "QTY_US", default=0),
                    )  # N
                    set_cell_safe(
                        ws,
                        row_idx,
                        15,
                        get_val(
                            row_data,
                            "QTY IN REPAIR",
                            "QTY_REPAIR",
                            "QTY REPAIR",
                            default=0,
                        ),
                    )  # O

                    set_cell_safe(
                        ws,
                        row_idx,
                        16,
                        get_val(
                            row_data,
                            "SHELF LIFE EXPIRATION",
                            "SHELF LIFE EXP",
                            "TOOL LIFE EXPIRATION",
                            default="",
                        ),
                    )  # P: Shelf Life Exp
                    set_cell_safe(
                        ws,
                        row_idx,
                        17,
                        get_val(row_data, "CONDITION", default="SV"),
                    )  # Q: Condition
                    set_cell_safe(
                        ws,
                        row_idx,
                        18,
                        get_val(row_data, "CATEGORY", default=""),
                    )  # R: Category
                    set_cell_safe(
                        ws,
                        row_idx,
                        19,
                        get_val(row_data, "OWNER", default=""),
                    )  # S: Owner
                    set_cell_safe(
                        ws,
                        row_idx,
                        20,
                        get_val(row_data, "UOM", default="EA"),
                    )  # T: UOM

                    # --- FORMULA EXCEL DINAMIS ---
                    set_cell_safe(
                        ws,
                        row_idx,
                        21,
                        f"=J{row_idx}+K{row_idx}+M{row_idx}+N{row_idx}",
                    )  # U: Qty eMRO
                    set_cell_safe(ws, row_idx, 22, f"=U{row_idx}")  # V: Qty Actual
                    set_cell_safe(
                        ws, row_idx, 23, f"=V{row_idx}-U{row_idx}"
                    )  # W: Diff
                    set_cell_safe(
                        ws,
                        row_idx,
                        24,
                        f'=IF(U{row_idx}=0,"BUG EMRO??/MISSING??",IF(V{row_idx}=0,"NOT FOUND",IF(V{row_idx}>U{row_idx},"SURPLUS",IF(V{row_idx}<U{row_idx},"MINUS","MATCHED"))))',
                    )  # X: Result
                    set_cell_safe(
                        ws,
                        row_idx,
                        25,
                        f'=IF(X{row_idx}="MATCHED","MATCHED","OPEN")',
                    )  # Y: Status

                    # Dikosongkan
                    set_cell_safe(ws, row_idx, 26, None)  # Z: Date
                    set_cell_safe(ws, row_idx, 27, None)  # AA: Auditor
                    set_cell_safe(ws, row_idx, 28, None)  # AB: Remark
                    set_cell_safe(ws, row_idx, 29, None)  # AC: Penyelesaian
                    set_cell_safe(ws, row_idx, 30, None)  # AD: Corrective Action

                    # --- PENUKARAN KOLOM PREV SO & PREV STATUS ---
                    prev_info = prev_so_map.get(
                        batch_num, {"prev_so": None, "prev_status": None}
                    )
                    set_cell_safe(
                        ws, row_idx, 31, prev_info["prev_so"]
                    )  # AE: Prev SO
                    set_cell_safe(
                        ws,
                        row_idx,
                        32,
                        get_val(row_data, "CAT", "CATEGORY", default=""),
                    )  # AF: CAT
                    set_cell_safe(
                        ws, row_idx, 33, prev_info["prev_status"]
                    )  # AG: Prev Status
                    set_cell_safe(ws, row_idx, 34, None)  # AH: Reason

            # -----------------------------------------------------
            # UPDATE SHEET SUMMARY
            # -----------------------------------------------------
            if "Summary" in wb.sheetnames:
                ws_sum = wb["Summary"]
                set_cell_safe(ws_sum, 3, 4, f": {station_input}")  # D3
                set_cell_safe(ws_sum, 4, 4, f": {periode_input}")  # D4

                start_sum_row = 11  # Data Summary dimulai di baris 11
                for i, sdata in enumerate(summary_data_inputs):
                    r_curr = start_sum_row + i
                    set_cell_safe(ws_sum, r_curr, 2, i + 1)  # B: NO
                    set_cell_safe(
                        ws_sum, r_curr, 3, sdata["division"]
                    )  # C: DIVISION
                    set_cell_safe(ws_sum, r_curr, 4, sdata["pic"])  # D: PIC
                    set_cell_safe(
                        ws_sum, r_curr, 5, sdata["loc_code"]
                    )  # E: LOC CODE
                    set_cell_safe(
                        ws_sum, r_curr, 6, sdata["loc_desc"]
                    )  # F: LOCATION DESCRIPTION

            # Save ke memory buffer
            output_buffer = io.BytesIO()
            wb.save(output_buffer)
            output_buffer.seek(0)

            st.success("✅ Dokumen berhasil diproses dan dirapikan!")

            st.download_button(
                label="📥 Download Hasil Excel Stock Opname",
                data=output_buffer,
                file_name=f"Worksheet_Stock_Opname_{station_input}_Updated.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
